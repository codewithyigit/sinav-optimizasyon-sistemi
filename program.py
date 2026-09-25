import http.server
import socketserver
import pandas as pd
import io
import email
import email.policy
import base64
import traceback
import sys

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


SERVER_PORT = 1903
EXCEL_FONT_NAME = 'Calibri'
COLOR_HEADER_BG = "4472C4"
COLOR_HEADER_TXT = "FFFFFF"
COLOR_COMMON_BG = "FFEB9C"  
COLOR_COMMON_TXT = "9C5700"
REQUIRED_ROOM_COLS = ['Oda', 'Kapasite', 'Zaman']


def apply_excel_formatting(ws):
    """
    Excel sayfasına stil ve renklendirme uygular.
    Başlıkları koyu mavi, ortak dersleri sarı yapar.
    """
   
    header_style = {
        'font': Font(bold=True, color=COLOR_HEADER_TXT, size=11, name=EXCEL_FONT_NAME),
        'fill': PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid"),
        'alignment': Alignment(horizontal="center", vertical="center", wrap_text=True)
    }
    
    common_fill = PatternFill(start_color=COLOR_COMMON_BG, end_color=COLOR_COMMON_BG, fill_type="solid")
    common_font = Font(color=COLOR_COMMON_TXT, name=EXCEL_FONT_NAME)
    
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), 
                         top=Side(style="thin"), bottom=Side(style="thin"))

    for cell in ws[1]:
        cell.font = header_style['font']
        cell.fill = header_style['fill']
        cell.alignment = header_style['alignment']
        cell.border = thin_border

    for row in ws.iter_rows(min_row=2):
 
        first_val = str(row[0].value).lower() if row[0].value else ""
        is_common_course = "ortak" in first_val

        for cell in row:
            cell.border = thin_border
            
            if is_common_course:
                cell.fill = common_fill
                cell.font = common_font
            
            if cell.value and len(str(cell.value)) > 50:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, col_cells in enumerate(ws.columns, 1):
        col_letter = get_column_letter(i)
        header_val = str(col_cells[0].value)
        
        if "Listesi" in header_val:
            ws.column_dimensions[col_letter].width = 50
        else:
            max_len = 0
            for cell in col_cells:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


class ExamScheduler:
    """Sınav dağıtım algoritmasını yöneten sınıf."""
    
    def __init__(self):
        self.results = []
        self.busy_map = {
            'students': {},
            'proctors': {},
            'rooms': {}
        }

    def _load_registrations(self, df_regs):
        course_map = df_regs.groupby('Ders Kodu')['Öğrenci No'].apply(list).to_dict()
        
        student_names = {}
        if 'Ad Soyad' in df_regs.columns:
            tmp = df_regs[['Öğrenci No', 'Ad Soyad']].drop_duplicates()
            student_names = pd.Series(tmp['Ad Soyad'].values, index=tmp['Öğrenci No']).to_dict()
            
        return course_map, student_names

    def _prepare_rooms(self, df_rooms):
        time_slots = df_rooms['Zaman'].drop_duplicates().tolist()
        room_map = {}
        
        for _, row in df_rooms.iterrows():
            t = row['Zaman']
            if t not in room_map: 
                room_map[t] = []
            
            room_map[t].append({
                "name": row['Oda'],
                "cap": int(row['Kapasite'])
            })
        return time_slots, room_map

    def run(self, df_rooms, df_exams, df_regs):
        self.results = [] 
        
        if not all(c in df_rooms.columns for c in REQUIRED_ROOM_COLS):
            return None, "HATA: Oda dosyasında sütun isimleri hatalı. (Oda, Kapasite, Zaman) gerekli."

        if 'Tip' not in df_exams.columns:
            df_exams['Tip'] = 'Bölüm'
        
        df_exams['Tip'] = df_exams['Tip'].fillna('Bölüm').astype(str)
        df_exams['priority'] = df_exams['Tip'].apply(lambda x: 0 if 'ortak' in x.lower() else 1)
        df_exams = df_exams.sort_values(by=['priority', 'Ders Kodu'])
        course_students, student_names = self._load_registrations(df_regs)
        time_slots, room_map = self._prepare_rooms(df_rooms)

        print(f"Toplam {len(df_exams)} sınav planlanacak.")

        for _, row in df_exams.iterrows():
            code = row['Ders Kodu']
            c_name = row.get('Ders Adı', '')
            proctor = row['Sorumlu Hoca']
            c_type = row['Tip']
            students = course_students.get(code, [])
            count = len(students)

            if count == 0:
                print(f"UYARI: {code} kodlu dersin öğrencisi yok atlanıyor.")
                self.results.append({
                    "Tip": c_type, "Ders Kodu": code, "Ders Adı": c_name, "Durum": "Hata: Öğrenci Yok"
                })
                continue

            is_placed = False
            
            for time in time_slots:
                if is_placed: break
             
                available_rooms = room_map.get(time, [])
                if not available_rooms: continue

                if self._check_conflict('proctors', proctor, time): continue
        
                student_conflict = False
                for sid in students:
                    if self._check_conflict('students', sid, time):
                        student_conflict = True
                        break
                if student_conflict: continue

                
                selected_room = None
                selected_cap = 0
                
                for r in available_rooms:
                    r_name = r['name']
                   
                    if self._check_conflict('rooms', r_name, time): continue
                    
                    if r['cap'] >= count:
                        selected_room = r_name
                        selected_cap = r['cap']
                        break 
                
                if selected_room:
                  
                    is_placed = True
                    self._register_busy('rooms', selected_room, time)
                    self._register_busy('proctors', proctor, time)
                    for sid in students:
                        self._register_busy('students', sid, time)

                   
                    s_list_str = [f"{sid} - {student_names.get(sid, 'Bilinmiyor')}" for sid in students]
                    
                    self.results.append({
                        "Tip": c_type,
                        "Ders Kodu": code,
                        "Ders Adı": c_name,
                        "Sınıf": selected_room,
                        "Saat": time,
                        "Öğrenci Sayısı": count,
                        "Doluluk": f"{count} / {selected_cap}",
                        "Gözetmen": proctor,
                        "Öğrenci Listesi": "\n".join(s_list_str),
                        "Durum": "Başarılı"
                    })

            if not is_placed:
                print(f"HATA: {code} dersi yerleştirilemedi (Yer/Zaman yok).")
                self.results.append({
                    "Tip": c_type,
                    "Ders Kodu": code,
                    "Ders Adı": c_name,
                    "Durum": "BAŞARISIZ: Uygun Aralık Bulunamadı"
                })

        return pd.DataFrame(self.results), None

    def _check_conflict(self, category, key, time):
        """Yardımcı fonksiyon: Belirtilen varlık o saatte meşgul mü?"""
        if key in self.busy_map[category]:
            return time in self.busy_map[category][key]
        return False

    def _register_busy(self, category, key, time):
        """Yardımcı fonksiyon: Varlığı o saate kilitle."""
        if key not in self.busy_map[category]:
            self.busy_map[category][key] = []
        self.busy_map[category][key].append(time)


class SinavSunucusu(http.server.SimpleHTTPRequestHandler):
    
    def get_upload_page(self):
        return """
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <title>Sınav Otomasyonu v1.0</title>
            <style>
                body { font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; padding: 20px; background: #f4f4f9; }
                .container { background: white; padding: 40px; border-radius: 8px; box-shadow: 0 4px 6px rgba(0,0,0,0.1); max-width: 500px; margin: 40px auto; }
                h2 { text-align:center; color: #333; margin-top: 0; }
                .form-group { margin-bottom: 20px; }
                label { display: block; font-weight: 600; margin-bottom: 8px; color: #555; }
                input[type="file"] { border: 1px solid #ddd; padding: 10px; width: 100%; box-sizing: border-box; border-radius: 4px; }
                button { background: #4472C4; color: white; padding: 12px; border: none; width: 100%; cursor: pointer; font-size: 16px; border-radius: 4px; transition: background 0.3s; }
                button:hover { background: #365a9e; }
                .info { font-size: 13px; color: #856404; background: #fff3cd; padding: 10px; border-radius: 4px; margin-bottom: 20px; text-align: center; }
            </style>
        </head>
        <body>
            <div class="container">
                <h2>Sınav Programlayıcı</h2>
                <div class="info">Lütfen Excel (.xlsx) dosyalarını yükleyin.</div>
                
                <form action="/" method="POST" enctype="multipart/form-data">
                    <div class="form-group">
                        <label>1. Sınav Listesi</label>
                        <input type="file" name="file_exams" accept=".xlsx" required>
                    </div>
                    <div class="form-group">
                        <label>2. Öğrenci Kayıtları</label>
                        <input type="file" name="file_regs" accept=".xlsx" required>
                    </div>
                    <div class="form-group">
                        <label>3. Oda ve Zaman Tablosu</label>
                        <input type="file" name="file_rooms" accept=".xlsx" required>
                    </div>
                    <button type="submit">Programı Oluştur</button>
                </form>
            </div>
        </body>
        </html>
        """

    def do_GET(self):
        if self.path == '/':
            self.send_response(200)
            self.send_header('Content-type', 'text/html; charset=utf-8')
            self.end_headers()
            self.wfile.write(self.get_upload_page().encode('utf-8'))

    def do_POST(self):
        try:
          
            content_length = int(self.headers['Content-Length'])
            body = self.rfile.read(content_length)
            
           
            headers = f"Content-Type: {self.headers['Content-Type']}\r\n"
            msg = email.message_from_bytes(headers.encode() + b'\r\n' + body, policy=email.policy.HTTP)

            files = {}
            for part in msg.iter_parts():
                if part.get_content_disposition() != 'form-data': continue
                field_name = part.get_param('name', header='content-disposition')
                files[field_name] = part.get_payload(decode=True)

            if 'file_exams' not in files or 'file_regs' not in files or 'file_rooms' not in files:
                raise ValueError("Eksik dosya yüklendi.")

            
            
            df_exams = pd.read_excel(io.BytesIO(files['file_exams']))
            df_regs = pd.read_excel(io.BytesIO(files['file_regs']))
            df_rooms = pd.read_excel(io.BytesIO(files['file_rooms']))

           
            scheduler = ExamScheduler()
            result_df, error_msg = scheduler.run(df_rooms, df_exams, df_regs)

            if error_msg:
                self.send_response(400)
                self.end_headers()
                self.wfile.write(f"Mantıksal Hata: {error_msg}".encode('utf-8'))
                return

            
            output_buffer = io.BytesIO()
            with pd.ExcelWriter(output_buffer, engine='openpyxl') as writer:
                
                result_df.to_excel(writer, index=False, sheet_name='Tüm Program')
                apply_excel_formatting(writer.sheets['Tüm Program'])
                
              
                df_ortak = result_df[result_df['Tip'].str.contains('Ortak', case=False, na=False)]
                if not df_ortak.empty:
                    df_ortak.to_excel(writer, index=False, sheet_name='Sadece Ortaklar')
                    apply_excel_formatting(writer.sheets['Sadece Ortaklar'])
                
                df_bolum = result_df[~result_df['Tip'].str.contains('Ortak', case=False, na=False)]
                if not df_bolum.empty:
                    df_bolum.to_excel(writer, index=False, sheet_name='Bölüm Dersleri')
                    apply_excel_formatting(writer.sheets['Bölüm Dersleri'])

           
            b64_data = base64.b64encode(output_buffer.getvalue()).decode()
            
            self.send_response(200)
            self.send_header('Content-type', 'text/html; charset=utf-8')
            self.end_headers()
            
           
            table_html = result_df.to_html(index=False, classes='preview-table', escape=False).replace('\\n', '<br>')
            
            response_html = f"""
            <html>
            <head>
                <meta charset="utf-8">
                <style>
                    body {{ padding:20px; font-family: sans-serif; background: #eee; }}
                    .card {{ background: white; padding: 20px; border-radius: 8px; box-shadow: 0 2px 4px rgba(0,0,0,0.1); }}
                    h2 {{ color: #333; margin-top:0; }}
                    table {{ width: 100%; border-collapse: collapse; font-size: 13px; margin-top: 15px; }}
                    th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; vertical-align: top; }}
                    th {{ background-color: #4472C4; color: white; }}
                    tr.ortak-row td {{ background-color: #fff3cd !important; }}
                    .btn {{ display: inline-block; padding: 10px 20px; color: white; text-decoration: none; border-radius: 4px; font-weight:bold; margin-right: 10px; }}
                    .btn-download {{ background: #28a745; }}
                    .btn-back {{ background: #6c757d; }}
                </style>
            </head>
            <body>
                <div class="card">
                    <h2>✅ İşlem Başarılı</h2>
                    <p>Program başarıyla oluşturuldu. Aşağıdaki butondan indirebilirsiniz.</p>
                    <a href="data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,{b64_data}" download="Sinav_Programi.xlsx" class="btn btn-download">📥 Excel İndir</a>
                    <a href="/" class="btn btn-back">Geri Dön</a>
                </div>
                
                <div class="card" style="margin-top: 20px;">
                    <h3>Önizleme</h3>
                    {table_html}
                </div>

                <script>
                    // Basit bir JS ile ortak dersleri vurgula
                    document.querySelectorAll('table tr').forEach(row => {{
                        if (row.innerText.toLowerCase().includes('ortak')) {{
                            row.classList.add('ortak-row');
                        }}
                    }});
                </script>
            </body>
            </html>
            """
            self.wfile.write(response_html.encode('utf-8'))

        except Exception as e:
            
            print("Sunucu tarafında hata oluştu:")
            traceback.print_exc()
            self.send_response(500)
            self.end_headers()
            self.wfile.write(f"Sunucu Hatasi: {str(e)}".encode('utf-8'))


class ReuseAddrServer(socketserver.TCPServer):
    allow_reuse_address = True

if __name__ == "__main__":
  
    print(f" Sınav Programlayıcı Başlatılıyor...")
    print(f" Adres: http://localhost:{SERVER_PORT}")
    print(f" Durdurmak için: CTRL+C")

    httpd = ReuseAddrServer(("", SERVER_PORT), SinavSunucusu)
    
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\nKapatma sinyali alındı. Çıkılıyor...")
        httpd.server_close()
        sys.exit(0)